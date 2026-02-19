#!/usr/bin/env python3
# =============================================================================
# 01_extract_osp_felodipine.py
# Extract felodipine oral PK profiles from OSP observed-data database
# =============================================================================
#
# Input:  OSP ObsDataPK_OSP.xlsx (shared via portfolio project 02)
# Output: data/observed/extracted/mean_profiles_felodipine_po.csv
#         data/observed/extracted/study_table_felodipine_po.csv
#         data/sources/provenance.md
#
# Usage:  python3 analysis/01_extract_osp_felodipine.py

import os
import sys
import csv
import math
from datetime import datetime

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl required. Install with: pip install openpyxl")
    sys.exit(1)

# --- Import setup -----------------------------------------------------------
import importlib.util
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location(
    "setup_00", os.path.join(_SCRIPT_DIR, "00_setup.py"))
_setup = importlib.util.module_from_spec(_spec)
sys.modules["setup_00"] = _setup
_spec.loader.exec_module(_setup)

PROJECT_ROOT = _setup.PROJECT_ROOT
DIR_OBSERVED_EXTRACTED = _setup.DIR_OBSERVED_EXTRACTED
DIR_SOURCES = _setup.DIR_SOURCES
OSP_XLSX = _setup.OSP_XLSX

# --- Constants --------------------------------------------------------------
MW_FELODIPINE = 384.26  # g/mol

UNIT_FACTORS = {
    "ng/mL": 1.0,
    "ng/ml": 1.0,
    "µg/L":  1.0,
    "ug/L":  1.0,
    "mg/L":  1000.0,
    "pg/mL": 0.001,
    "nmol/l": MW_FELODIPINE / 1000.0,
    "nmol/L": MW_FELODIPINE / 1000.0,
    "ng/L":  0.001,
}

# DDI perpetrators to exclude (keep only control/baseline arms)
DDI_PERPETRATORS = {
    "grapefruit", "ketoconazole", "itraconazole", "erythromycin",
    "cimetidine", "rifampin", "rifampicin", "st john",
}

# --- Logging ----------------------------------------------------------------
log_lines = []


def log(msg):
    ts = datetime.now().strftime("%H:%M:%S")
    line = f"[{ts}] {msg}"
    print(line)
    log_lines.append(line)


# --- Load workbook ----------------------------------------------------------
if not os.path.exists(OSP_XLSX):
    print(f"ERROR: OSP data file not found at {OSP_XLSX}")
    print("Please ensure the file exists (from project 02-POPPK-midazolam).")
    sys.exit(1)

log(f"Loading {OSP_XLSX}")
wb = openpyxl.load_workbook(OSP_XLSX, data_only=True)

# --- Parse Studies sheet ----------------------------------------------------
log("Parsing Studies sheet...")
ws_studies = wb["Studies"]
studies = {}
for row_idx in range(3, ws_studies.max_row + 1):
    sid = ws_studies.cell(row=row_idx, column=1).value
    if sid is None:
        continue
    studies[sid] = {
        "study_id": sid,
        "study": ws_studies.cell(row=row_idx, column=2).value,
        "reference": ws_studies.cell(row=row_idx, column=3).value,
        "grouping": ws_studies.cell(row=row_idx, column=4).value,
        "compound": ws_studies.cell(row=row_idx, column=5).value,
        "compartment": ws_studies.cell(row=row_idx, column=6).value,
        "data_type": ws_studies.cell(row=row_idx, column=7).value,
        "source": ws_studies.cell(row=row_idx, column=8).value,
        "dose_str": ws_studies.cell(row=row_idx, column=10).value,
        "dose_unit": ws_studies.cell(row=row_idx, column=12).value,
        "route": ws_studies.cell(row=row_idx, column=13).value,
        "fasted_fed": ws_studies.cell(row=row_idx, column=22).value,
        "species": ws_studies.cell(row=row_idx, column=26).value,
        "n": ws_studies.cell(row=row_idx, column=27).value,
    }

log(f"  Total studies parsed: {len(studies)}")

# --- Filter: Felodipine, PO, Human -----------------------------------------
felo_study_ids = set()
for sid, s in studies.items():
    compound = str(s["compound"] or "").lower()
    route = str(s["route"] or "").lower()
    species = str(s["species"] or "").lower()
    if "felodipine" in compound and "po" in route and "human" in species:
        felo_study_ids.add(sid)

log(f"  Felodipine PO Human studies: {len(felo_study_ids)}")

# --- Classify: control vs DDI arms -----------------------------------------
control_ids = set()
ddi_ids = set()

for sid in felo_study_ids:
    s = studies[sid]
    grouping = str(s.get("grouping") or "").lower()
    reference = str(s.get("reference") or "").lower()
    study_name = str(s.get("study") or "").lower()

    # Check if this arm involves a DDI perpetrator
    combined = f"{grouping} {reference} {study_name}"
    is_ddi = any(perp in combined for perp in DDI_PERPETRATORS)

    # Explicit control labels
    is_control = any(tag in grouping for tag in
                     ["control", "water", "baseline", "placebo"])

    if is_ddi and not is_control:
        ddi_ids.add(sid)
    else:
        control_ids.add(sid)

log(f"  Control/baseline arms: {len(control_ids)}")
log(f"  DDI treatment arms (excluded): {len(ddi_ids)}")

# --- Parse dose strings -----------------------------------------------------
def parse_dose_mg(dose_str, dose_unit):
    if dose_str is None:
        return None
    ds = str(dose_str).strip()
    du = str(dose_unit or "mg").strip().lower()
    if "mg/kg" in du or "mg/kg" in ds.lower():
        return None
    if "(" in ds:
        ds = ds.split("(")[0].strip()
    ds = ds.replace("mg", "").strip()
    if "-" in ds and not ds.startswith("-"):
        parts = ds.split("-")
        try:
            return float(parts[0].strip())
        except ValueError:
            return None
    try:
        val = float(ds)
        if "µg" in str(dose_unit or "") or "ug" in str(dose_unit or ""):
            val = val / 1000.0
        return val
    except ValueError:
        return None


# --- Parse PK-Profiles sheet ------------------------------------------------
log("Parsing PK-Profiles sheet...")
ws_profiles = wb["PK-Profiles"]

profiles_raw = []
for row_idx in range(2, ws_profiles.max_row + 1):
    sid = ws_profiles.cell(row=row_idx, column=1).value
    if sid not in felo_study_ids:
        continue

    analyte = ws_profiles.cell(row=row_idx, column=5).value
    if not analyte or "felodipine" not in str(analyte).lower():
        continue

    compartment = ws_profiles.cell(row=row_idx, column=6).value

    profiles_raw.append({
        "study_id": sid,
        "study_label": ws_profiles.cell(row=row_idx, column=2).value,
        "reference": ws_profiles.cell(row=row_idx, column=3).value,
        "grouping": ws_profiles.cell(row=row_idx, column=4).value,
        "analyte": analyte,
        "compartment": str(compartment or ""),
        "time": ws_profiles.cell(row=row_idx, column=7).value,
        "time_unit": ws_profiles.cell(row=row_idx, column=8).value,
        "avg": ws_profiles.cell(row=row_idx, column=9).value,
        "avg_unit": ws_profiles.cell(row=row_idx, column=10).value,
        "avg_type": ws_profiles.cell(row=row_idx, column=11).value,
        "var": ws_profiles.cell(row=row_idx, column=12).value,
        "var_unit": ws_profiles.cell(row=row_idx, column=13).value,
        "var_type": ws_profiles.cell(row=row_idx, column=14).value,
    })

log(f"  Felodipine PO profile rows (all compartments): {len(profiles_raw)}")

# --- Filter to Plasma only --------------------------------------------------
profiles_plasma = [p for p in profiles_raw
                   if "plasma" in p["compartment"].lower()]
log(f"  Plasma only: {len(profiles_plasma)}")

excluded_compartments = set(p["compartment"] for p in profiles_raw
                           if "plasma" not in p["compartment"].lower())
if excluded_compartments:
    log(f"  Excluded compartments: {excluded_compartments}")

# --- Normalize units to ng/mL -----------------------------------------------
skipped_units = set()
profiles_norm = []
for p in profiles_plasma:
    unit = str(p["avg_unit"] or "").strip()
    factor = UNIT_FACTORS.get(unit)

    if factor is None:
        skipped_units.add(unit)
        continue

    avg_val = p["avg"]
    if avg_val is None or not isinstance(avg_val, (int, float)):
        continue

    conc_ngml = float(avg_val) * factor

    # Convert variability to SD in ng/mL
    sd_ngml = None
    var_val = p["var"]
    var_type = str(p["var_type"] or "").strip().lower()
    var_unit = str(p["var_unit"] or "").strip()
    var_factor = UNIT_FACTORS.get(var_unit, factor)

    study_info = studies.get(p["study_id"], {})
    n = study_info.get("n")
    if n is not None:
        try:
            n = int(n)
        except (ValueError, TypeError):
            n = None

    if var_val is not None and isinstance(var_val, (int, float)):
        var_converted = float(var_val) * var_factor
        if "sd" in var_type:
            sd_ngml = var_converted
        elif "sem" in var_type:
            if n and n > 0:
                sd_ngml = var_converted * math.sqrt(n)

    dose_mg = parse_dose_mg(study_info.get("dose_str"),
                            study_info.get("dose_unit"))

    time_val = p["time"]
    if time_val is None or not isinstance(time_val, (int, float)):
        continue

    # Classify fasted vs fed
    grouping = str(p.get("grouping") or "").lower()
    fasted_fed_col = str(study_info.get("fasted_fed") or "").lower()
    if any(tag in grouping or tag in fasted_fed_col for tag in ["fed", "food", "meal"]):
        prandial = "fed"
    elif any(tag in grouping or tag in fasted_fed_col for tag in ["fasted", "fasting"]):
        prandial = "fasted"
    else:
        prandial = "unknown"

    # Control arm flag
    is_control = p["study_id"] in control_ids

    profiles_norm.append({
        "study_id": p["study_id"],
        "study_label": p["study_label"],
        "grouping": p["grouping"] or "",
        "time_h": float(time_val),
        "conc_ngml": round(conc_ngml, 4),
        "sd_ngml": round(sd_ngml, 4) if sd_ngml is not None else "",
        "n": n or "",
        "dose_mg": dose_mg or "",
        "data_type": study_info.get("data_type", ""),
        "prandial": prandial,
        "is_control": is_control,
    })

if skipped_units:
    log(f"  Skipped non-concentration units: {skipped_units}")
log(f"  Normalized plasma profiles: {len(profiles_norm)}")

# --- Separate control and DDI rows ------------------------------------------
control_rows = [p for p in profiles_norm if p["is_control"]]
ddi_rows = [p for p in profiles_norm if not p["is_control"]]
log(f"  Control arm data points: {len(control_rows)}")
log(f"  DDI arm data points (excluded from qualification): {len(ddi_rows)}")

# --- Write mean_profiles_felodipine_po.csv ----------------------------------
mean_csv_path = os.path.join(DIR_OBSERVED_EXTRACTED,
                             "mean_profiles_felodipine_po.csv")
fieldnames = ["study_id", "study_label", "grouping", "time_h", "conc_ngml",
              "sd_ngml", "n", "dose_mg", "prandial", "is_control"]
with open(mean_csv_path, "w", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=fieldnames)
    writer.writeheader()
    for row in sorted(profiles_norm,
                      key=lambda x: (x["study_id"], x["time_h"])):
        writer.writerow({k: row[k] for k in fieldnames})

log(f"  Wrote {mean_csv_path} ({len(profiles_norm)} rows)")

# --- Write study_table_felodipine_po.csv ------------------------------------
study_table_path = os.path.join(DIR_OBSERVED_EXTRACTED,
                                "study_table_felodipine_po.csv")
study_agg = {}
for row in profiles_norm:
    sid = row["study_id"]
    if sid not in study_agg:
        info = studies.get(sid, {})
        study_agg[sid] = {
            "study_id": sid,
            "study_label": row["study_label"],
            "reference": info.get("reference", ""),
            "grouping": row["grouping"],
            "n": info.get("n", ""),
            "dose_mg": row["dose_mg"],
            "n_timepoints": 0,
            "time_min_h": float("inf"),
            "time_max_h": float("-inf"),
            "prandial": row["prandial"],
            "is_control": row["is_control"],
        }
    study_agg[sid]["n_timepoints"] += 1
    t = row["time_h"]
    study_agg[sid]["time_min_h"] = min(study_agg[sid]["time_min_h"], t)
    study_agg[sid]["time_max_h"] = max(study_agg[sid]["time_max_h"], t)

study_table_fields = ["study_id", "study_label", "reference", "grouping", "n",
                      "dose_mg", "n_timepoints", "time_range_h", "prandial",
                      "is_control"]
with open(study_table_path, "w", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=study_table_fields)
    writer.writeheader()
    for sid in sorted(study_agg.keys()):
        s = study_agg[sid]
        s["time_range_h"] = f"{s['time_min_h']:.1f}-{s['time_max_h']:.1f}"
        writer.writerow({k: s[k] for k in study_table_fields})

log(f"  Wrote {study_table_path} ({len(study_agg)} studies)")

# --- Write provenance.md ----------------------------------------------------
n_unique_studies = len(set(r["study_id"] for r in profiles_norm))
n_control = len(set(r["study_id"] for r in control_rows))
provenance_path = os.path.join(DIR_SOURCES, "provenance.md")
with open(provenance_path, "w") as f:
    f.write("# Data Provenance — Felodipine PK Profiles\n\n")
    f.write("## Source\n\n")
    f.write("- **Database:** Open-Systems-Pharmacology/Database-for-observed-data\n")
    f.write("- **File:** ObsDataPK_OSP.xlsx\n")
    f.write(f"- **Extraction date:** {datetime.now().strftime('%Y-%m-%d')}\n")
    f.write("- **URL:** https://github.com/Open-Systems-Pharmacology/Database-for-observed-data\n\n")
    f.write("## Extraction Pipeline\n\n")
    f.write("Script: `analysis/01_extract_osp_felodipine.py`\n\n")
    f.write("### Filtering Steps\n\n")
    f.write("| Step | Filter | Rows Remaining |\n")
    f.write("|------|--------|----------------|\n")
    f.write(f"| 0 | All PK-Profiles rows | {ws_profiles.max_row - 1} |\n")
    f.write(f"| 1 | Study ID in felodipine PO human studies | {len(profiles_raw)} |\n")
    f.write(f"| 2 | Compartment = Plasma | {len(profiles_plasma)} |\n")
    f.write(f"| 3 | Valid concentration units (ng/mL convertible) | {len(profiles_norm)} |\n")
    f.write(f"| 4 | Control/baseline arms only | {len(control_rows)} |\n\n")
    f.write("### Summary\n\n")
    f.write(f"- Unique studies with extracted profiles: {n_unique_studies}\n")
    f.write(f"- Control arm studies (qualification set): {n_control}\n")
    f.write(f"- Total data points: {len(profiles_norm)}\n")
    f.write(f"- Control arm data points: {len(control_rows)}\n")
    if skipped_units:
        f.write(f"- Skipped non-concentration units: {', '.join(sorted(skipped_units))}\n")
    if excluded_compartments:
        f.write(f"- Excluded compartments: {', '.join(sorted(excluded_compartments))}\n")
    f.write("\n### DDI Arms Excluded\n\n")
    f.write("Perpetrator co-administration arms were excluded from the qualification set:\n")
    for sid in sorted(ddi_ids):
        info = studies.get(sid, {})
        f.write(f"- Study {sid}: {info.get('study', '')} — {info.get('grouping', '')}\n")

log(f"  Wrote {provenance_path}")
log(f"\nExtraction complete: {len(profiles_norm)} data points from {n_unique_studies} studies.")
